"""Create the single-column Word CV from lib/cv.ts data exported by export-cv.mjs.

Usage: python scripts/export-cv-docx.py /path/to/cv-data.json
Requires python-docx. Font source/license: ../Cem Bilen CV 2026 HTML/fonts.
"""
from __future__ import annotations

import json
from pathlib import Path
import re
import sys
import uuid
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "public/cv/Cem Bilen CV 2026.docx"
FONTS = ROOT / "Cem Bilen CV 2026 HTML/fonts"
DATA = json.loads(Path(sys.argv[1]).read_text())

def plain(value):
    return re.sub(r"[\u2013\u2014\u2212]", "-", value)


def font(style, size, bold=False):
    style.font.name = "Inter"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string("262626")
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn(f"w:{kind}"), "Inter")
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = Pt(13.02)


def embed_fonts(filename):
    # ECMA-376 obfuscated OpenType parts allow the editable document to travel
    # with its OFL-licensed Inter fonts without relying on recipient installs.
    with ZipFile(filename) as src:
        parts = {name: src.read(name) for name in src.namelist()}
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    relns = "http://schemas.openxmlformats.org/package/2006/relationships"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ct = "http://schemas.openxmlformats.org/package/2006/content-types"
    table = etree.fromstring(parts["word/fontTable.xml"])
    inter = next((el for el in table if el.get(f"{{{w}}}name") == "Inter"), None)
    if inter is None:
        inter = etree.SubElement(table, f"{{{w}}}font", {f"{{{w}}}name": "Inter"})
        etree.SubElement(inter, f"{{{w}}}family", {f"{{{w}}}val": "swiss"})
        etree.SubElement(inter, f"{{{w}}}pitch", {f"{{{w}}}val": "variable"})
    relpath = "word/_rels/fontTable.xml.rels"
    rels = etree.fromstring(parts[relpath]) if relpath in parts else etree.Element(f"{{{relns}}}Relationships", nsmap={None: relns})
    for index, label in enumerate(("Regular", "Bold"), 1):
        key = uuid.uuid5(uuid.NAMESPACE_URL, "https://cxbilen.com/cv/Inter-" + label)
        key_bytes = bytes.fromhex(key.hex)[::-1]
        data = bytearray((FONTS / f"Inter-{label}.ttf").read_bytes())
        for i in range(32):
            data[i] ^= key_bytes[i % 16]
        file_name = f"Inter-{label}.odttf"
        parts[f"word/fonts/{file_name}"] = bytes(data)
        rid = f"rIdInter{index}"
        etree.SubElement(rels, f"{{{relns}}}Relationship", Id=rid,
                         Type=r + "/font", Target="fonts/" + file_name)
        etree.SubElement(inter, f"{{{w}}}embed{label}", {
            f"{{{r}}}id": rid, f"{{{w}}}fontKey": "{" + str(key).upper() + "}",
            f"{{{w}}}subsetted": "false",
        })
    content = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(content, f"{{{ct}}}Default", Extension="odttf", ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont")
    settings = etree.fromstring(parts["word/settings.xml"])
    etree.SubElement(settings, f"{{{w}}}embedTrueTypeFonts")
    for name, value in [("word/fontTable.xml", table), (relpath, rels), ("[Content_Types].xml", content), ("word/settings.xml", settings)]:
        parts[name] = etree.tostring(value, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(filename, "w", ZIP_DEFLATED) as dest:
        for name, data in parts.items():
            dest.writestr(name, data)


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(12)
font(doc.styles["Normal"], 10.5)
font(doc.styles["Title"], 25, True)
doc.styles["Title"].paragraph_format.line_spacing = Pt(27.5)
doc.styles["Title"].paragraph_format.space_after = Pt(3)
font(doc.styles["Heading 1"], 11.5, True)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(9)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
doc.styles["Heading 1"].paragraph_format.keep_with_next = True
font(doc.styles["Heading 2"], 10.5, True)
doc.styles["Heading 2"].paragraph_format.keep_with_next = True
# No tables, text boxes, images, headers, or footers: every item is a body paragraph.
# Clear built-in Word theme font/color references and title-rule residue.
# These can override explicit fonts or introduce a blue line in LibreOffice.
for element in list(doc.styles.element.iter()):
    if element.tag == qn("w:rFonts"):
        for attr in list(element.attrib):
            if "theme" in attr.lower():
                del element.attrib[attr]
    if element.tag == qn("w:color"):
        for attr in list(element.attrib):
            if "theme" in attr.lower():
                del element.attrib[attr]
    if element.tag == qn("w:pBdr"):
        element.getparent().remove(element)
doc.core_properties.title = "Cem Bilen Software Engineer CV"
doc.core_properties.author = DATA["name"]
doc.core_properties.subject = "Full stack and AI native software engineering"
doc.core_properties.keywords = "Software Engineer, Full-stack, AI-native"


def p(text="", style=None, before=None, after=None):
    paragraph = doc.add_paragraph(plain(text), style=style)
    if before is not None:
        paragraph.paragraph_format.space_before = Pt(before)
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    return paragraph


def labelled(label, value, before=0):
    paragraph = p(before=before)
    paragraph.add_run(plain(label + ": ")).bold = True
    paragraph.add_run(plain(value))
    return paragraph


p(DATA["name"], "Title")
role = p(f'{DATA["title"]} | {DATA["subtitle"]}', after=6)
role.runs[0].font.size = Pt(11.5)
contact = DATA["contact"]
p(f'{contact["email"]} | {contact["phone"]} | {contact["location"]}')
social = {item["label"]: item["href"] for item in DATA["socials"]}
p(f'Portfolio: {social["Portfolio"]} | GitHub: {social["GitHub"]}')
p(f'LinkedIn: {social["LinkedIn"]}')
p("Summary", "Heading 1")
p(DATA["about"])
p("Technical Skills", "Heading 1")
labelled("Technologies", ", ".join(DATA["tools"]) + ".")
labelled("Capabilities", ", ".join(DATA["skills"]) + ".", before=3)
p("Projects", "Heading 1")
for i, item in enumerate(DATA["selectedWork"]):
    labelled(item["name"], item["desc"], before=4 if i else 0)
p("Experience", "Heading 1")
for i, item in enumerate(DATA["experience"]):
    title = p(style="Heading 2", before=5 if i else 0)
    title.add_run(plain(f'{item["role"]} | {item["company"]} ')).bold = True
    period = title.add_run(plain(f'| {item["period"]}'))
    period.bold = False
    period.font.color.rgb = RGBColor.from_string("686868")
    title.paragraph_format.keep_with_next = bool(item["body"])
    if item["body"]:
        p(item["body"])
p("Education", "Heading 1")
for i, item in enumerate(DATA["education"]):
    title = p(style="Heading 2", before=5 if i else 0)
    title.add_run(plain(f'{item["title"]} | {item["org"]} ')).bold = True
    period = title.add_run(plain(f'| {item["period"]}'))
    period.bold = False
    period.font.color.rgb = RGBColor.from_string("686868")
    title.paragraph_format.keep_with_next = bool(item["body"])
    if item["body"]:
        p(item["body"])
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Inter"
        rf = run._element.get_or_add_rPr().get_or_add_rFonts()
        for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn(f"w:{kind}"), "Inter")
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
embed_fonts(OUT)
print(OUT)
